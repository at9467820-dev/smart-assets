"""
Excel Report Generation using openpyxl.
"""
from io import BytesIO
from datetime import date
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

NAVY_HEX = "12203D"
BRASS_HEX = "B98A2E"
LIGHT_HEX = "EEF0F4"
WHITE_HEX = "FFFFFF"


def _apply_header_style(cell):
    cell.font = Font(name="Calibri", bold=True, color=WHITE_HEX, size=10)
    cell.fill = PatternFill("solid", fgColor=NAVY_HEX)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(
        bottom=Side(style="thin", color="DEE1E8"),
        right=Side(style="thin", color="DEE1E8"),
    )


def _apply_cell_style(cell, row_idx):
    cell.font = Font(name="Calibri", size=9)
    fill_color = WHITE_HEX if row_idx % 2 == 0 else LIGHT_HEX
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.alignment = Alignment(vertical="center", wrap_text=False)
    cell.border = Border(
        bottom=Side(style="thin", color="DEE1E8"),
        right=Side(style="thin", color="DEE1E8"),
    )


def _add_title_row(ws, title, col_count):
    ws.insert_rows(1, 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    title_cell = ws.cell(row=1, column=1, value="AssetPulse — " + title)
    title_cell.font = Font(name="Calibri", bold=True, size=14, color=NAVY_HEX)
    title_cell.fill = PatternFill("solid", fgColor=LIGHT_HEX)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    sub_cell = ws.cell(row=2, column=1, value=f"Generated: {date.today().strftime('%d %B %Y')}")
    sub_cell.font = Font(name="Calibri", size=9, color="5B6273")
    ws.row_dimensions[2].height = 16


def generate_inventory_excel(assets) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory"

    headers = ["Asset Tag", "Name", "Category", "Department", "Location",
               "Status", "Purchase Date", "Cost (₹)", "Warranty Expiry",
               "Condition (1-10)", "Serial No.", "Manufacturer"]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        _apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    for row_idx, a in enumerate(assets, 2):
        row_data = [
            a.asset_tag, a.name, a.category,
            a.department.name if a.department else "",
            a.location or "",
            a.status_label,
            a.purchase_date.strftime("%d/%m/%Y") if a.purchase_date else "",
            float(a.purchase_cost or 0),
            a.warranty_expiry.strftime("%d/%m/%Y") if a.warranty_expiry else "",
            a.condition_rating or "",
            a.serial_number or "",
            a.manufacturer or "",
        ]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            _apply_cell_style(cell, row_idx)

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _add_title_row(ws, "Inventory Report", len(headers))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_maintenance_excel(logs) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Maintenance"

    headers = ["#", "Asset Tag", "Asset Name", "Issue Description", "Status",
               "Priority", "Requested By", "Assigned To", "Request Date",
               "Scheduled Date", "Completed Date", "Est. Cost (₹)", "Actual Cost (₹)", "Next Due Date"]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        _apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

    for row_idx, log in enumerate(logs, 2):
        row_data = [
            row_idx - 1,
            log.asset.asset_tag if log.asset else "",
            log.asset.name if log.asset else "",
            log.issue_description or "",
            log.status.replace("_", " ").title() if log.status else "",
            log.priority.title() if log.priority else "",
            log.requested_by.full_name if log.requested_by else "",
            log.assigned_to.full_name if log.assigned_to else "",
            log.request_date.strftime("%d/%m/%Y") if log.request_date else "",
            log.scheduled_date.strftime("%d/%m/%Y") if log.scheduled_date else "",
            log.completed_date.strftime("%d/%m/%Y") if log.completed_date else "",
            float(log.estimated_cost or 0),
            float(log.actual_cost or 0),
            log.next_due_date.strftime("%d/%m/%Y") if log.next_due_date else "",
        ]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            _apply_cell_style(cell, row_idx)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _add_title_row(ws, "Maintenance Report", len(headers))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_budget_excel(estimates) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"

    headers = ["Department", "Fiscal Year", "Category", "Est. Maintenance (₹)",
               "Est. Replacement (₹)", "Total Estimated (₹)", "Actual Spent (₹)", "Variance (₹)"]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        _apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col_idx)].width = 22

    for row_idx, e in enumerate(estimates, 2):
        total = float(e.estimated_maintenance_cost or 0) + float(e.estimated_replacement_cost or 0)
        actual = float(e.actual_spent or 0)
        variance = actual - total
        row_data = [
            e.department.name if e.department else "",
            e.fiscal_year,
            e.category or "All",
            float(e.estimated_maintenance_cost or 0),
            float(e.estimated_replacement_cost or 0),
            total,
            actual,
            variance,
        ]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            _apply_cell_style(cell, row_idx)
            if col_idx >= 4:
                cell.number_format = '#,##0.00'

    ws.freeze_panes = "A2"
    _add_title_row(ws, "Budget Report", len(headers))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
